"""Constrained DOCX SOP and XLSX RFQ adapters for approved local libraries."""
from io import BytesIO
from zipfile import ZipFile, ZIP_STORED, ZIP_DEFLATED
from hashlib import sha256
from decimal import Decimal, InvalidOperation


def preflight(data):
    if not isinstance(data, bytes) or not 1 <= len(data) <= 5_000_000:
        raise ValueError("package size")
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        names = [item.filename for item in entries]
        if len(entries) > 500 or len(names) != len(set(names)):
            raise ValueError("package entries")
        if sum(item.file_size for item in entries) > 20_000_000:
            raise ValueError("expanded package size")
        for item in entries:
            name = item.filename
            if (name.startswith("/") or ".." in name.split("/") or "\\" in name
                    or item.flag_bits & 1 or item.compress_type not in (ZIP_STORED, ZIP_DEFLATED)
                    or item.file_size > 2_000_000
                    or item.file_size > max(1, item.compress_size) * 100
                    or any(part in name.lower() for part in ("vbaproject", "embeddings/", "externallinks/"))):
                raise ValueError("unsupported package member")
            if name.endswith((".xml", ".rels")):
                with archive.open(item) as member:
                    raw = member.read(2_000_001)
                if len(raw) > 2_000_000:
                    raise ValueError("XML member bound")
                text = raw.decode("utf-8-sig")
                if "\x00" in text or "<!DOCTYPE" in text.upper() or "<!ENTITY" in text.upper():
                    raise ValueError("DTD/entity or unsupported encoding")
                if name.endswith(".rels"):
                    from defusedxml.ElementTree import fromstring
                    root = fromstring(text, forbid_dtd=True, forbid_entities=True, forbid_external=True)
                    if any(child.get("TargetMode") == "External" for child in root):
                        raise ValueError("external relationships unsupported")


def scope(data, source_id, tenant, readers):
    if (not isinstance(source_id, str) or not source_id or not isinstance(tenant, str) or not tenant
            or not isinstance(readers, tuple) or not readers
            or any(not isinstance(r, str) or not r for r in readers)):
        raise ValueError("verified source permissions required")
    preflight(data)
    return {"source_id": source_id, "version": sha256(data).hexdigest(),
            "tenant": tenant, "readers": readers}


def docx_sop(data, *, source_id, tenant, readers):
    from docx import Document
    metadata = scope(data, source_id, tenant, readers)
    document = Document(BytesIO(data))
    unsupported = {"ins", "del", "commentRangeStart", "commentReference", "footnoteReference",
                   "endnoteReference", "drawing", "pict", "hyperlink", "fldChar", "instrText"}
    if any(node.tag.rsplit("}", 1)[-1] in unsupported for node in document.element.iter()):
        raise ValueError("document feature requires richer extraction")
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    blocks = []
    # Keep top-level document order, including tables between paragraphs.
    for index, child in enumerate(document.element.body):
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            blocks.append({**metadata, "location": f"body/{index}", "kind": "paragraph",
                           "style": paragraph.style.name, "text": paragraph.text})
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            if any(node.tag.endswith(("}gridSpan", "}vMerge", "}tbl"))
                   for cell in table._tbl.iterchildren() for node in cell.iter()
                   if node is not table._tbl):
                # Nested/merged tables need the dedicated structural adapter.
                raise ValueError("complex table requires review")
            blocks.append({**metadata, "location": f"body/{index}", "kind": "table",
                           "rows": [[cell.text for cell in row.cells] for row in table.rows]})
        elif not child.tag.endswith("}sectPr"):
            raise ValueError("unsupported document block")
        if len(blocks) > 500:
            raise ValueError("too many blocks")
    if document.inline_shapes or any(
            part.tables or any(p.text for p in part.paragraphs)
            for section in document.sections
            for part in (section.header, section.footer, section.first_page_header,
                         section.first_page_footer, section.even_page_header, section.even_page_footer)):
        raise ValueError("image/header/footer requires richer extraction")
    return blocks


def xlsx_rfqs(data, *, source_id, tenant, readers):
    from openpyxl import load_workbook
    metadata = scope(data, source_id, tenant, readers)
    workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("one-sheet export required")
        sheet = workbook.worksheets[0]
        if (sheet.max_row > 1001 or sheet.max_column != 3 or sheet.merged_cells.ranges
                or sheet.sheet_state != "visible"):
            raise ValueError("sheet shape or visibility")
        if [sheet.cell(1, col).value for col in range(1, 4)] != ["rfq_id", "currency", "amount"]:
            raise ValueError("header schema")
        rows, seen = [], set()
        for index in range(2, sheet.max_row + 1):
            cells = [sheet.cell(index, col) for col in range(1, 4)]
            if any(c.data_type in {"f", "e"} or c.hyperlink for c in cells) or sheet.row_dimensions[index].hidden:
                raise ValueError("formula, error, link or hidden row")
            rfq, currency, raw = [c.value for c in cells]
            if (not isinstance(rfq, str) or not rfq.strip() or rfq in seen
                    or currency not in {"USD", "GBP", "EUR"} or raw is None or isinstance(raw, bool)):
                raise ValueError("invalid RFQ fields")
            try:
                amount = Decimal(str(raw))
            except InvalidOperation as exc:
                raise ValueError("invalid amount") from exc
            if not amount.is_finite() or not 0 < amount <= Decimal("1000000000000"):
                raise ValueError("amount range")
            seen.add(rfq)
            rows.append({**metadata, "rfq_id": rfq, "currency": currency, "amount": str(amount),
                         "cells": {name: f"{sheet.title}!{cell.coordinate}" for name, cell in
                                   zip(("rfq_id", "currency", "amount"), cells)}})
        return rows
    finally:
        workbook.close()
