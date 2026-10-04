from io import BytesIO
from zipfile import ZipFile
import unittest
from ingest import preflight, docx_sop, xlsx_rfqs


SCOPE = dict(source_id="synthetic", tenant="a", readers=("reviewer",))


class OfficeTests(unittest.TestCase):
    def test_word_order_and_permissions(self):
        from docx import Document
        doc = Document()
        doc.add_heading("Approve RFQ", level=1)
        doc.add_paragraph("Check the amount.")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Owner", "Reviewer"
        stream = BytesIO()
        doc.save(stream)
        blocks = docx_sop(stream.getvalue(), **SCOPE)
        self.assertEqual([b["kind"] for b in blocks], ["paragraph", "paragraph", "table"])
        self.assertEqual(blocks[0]["style"], "Heading 1")
        self.assertTrue(all(b["readers"] == ("reviewer",) for b in blocks))

    def test_excel_schema_missing_and_formula(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(["rfq_id", "currency", "amount"])
        sheet.append(["r1", "USD", "12.50"])
        def data():
            stream = BytesIO()
            book.save(stream)
            return stream.getvalue()
        self.assertEqual(xlsx_rfqs(data(), **SCOPE)[0]["amount"], "12.50")
        for bad in (None, "=1+1", "NaN"):
            sheet["C2"] = bad
            with self.assertRaises(ValueError):
                xlsx_rfqs(data(), **SCOPE)
        book.close()

    def test_malformed_and_external_package(self):
        from zipfile import BadZipFile
        with self.assertRaises(BadZipFile):
            preflight(b"not a package")
        stream = BytesIO()
        with ZipFile(stream, "w") as archive:
            archive.writestr("x.rels", '<Relationships><Relationship TargetMode="External" Target="https://example.invalid"/></Relationships>')
        with self.assertRaises(ValueError):
            preflight(stream.getvalue())


if __name__ == "__main__":
    unittest.main()
